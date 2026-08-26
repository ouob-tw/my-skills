import subprocess
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def test_statistics_title_uses_statement_period(tmp_path):
    html_path = tmp_path / "2026-06.html"
    html_path.write_text(
        """
        <html><body>
          <div>對帳單期間: 2026/06/01~2026/06/30</div>
          <section id="tab2">
            <div class="chart-section">$100</div>
            <div class="table-scroll"><table>
              <tr><th>日期</th><th>交易說明</th><th>金額</th><th>餘額</th></tr>
              <tr><td>06/01</td><td>存入</td><td>$100</td><td>$100</td></tr>
            </table></div>
          </section>
          <section id="tab4">
            <div class="chart-section">消費總金額 75</div>
            <div class="table-scroll"><table>
              <tr><th> 消費日 / 入帳日 </th><th>商戶</th><th>金額</th><th>外幣消費金額/換匯日</th></tr>
              <tr><td>2026.06.02/2026.06.03</td><td>商店</td><td>$50</td><td>USD 1.50/2026.06.03</td></tr>
              <tr><td>2026.06.04</td><td>另一商店</td><td>$25</td><td>-</td></tr>
            </table></div>
          </section>
        </body></html>
        """,
        encoding="utf-8",
    )

    script = Path(__file__).parents[1] / "scripts" / "gen_linebank_xlsx.py"
    output_dir = tmp_path / "output"
    subprocess.run(
        [
            "uv",
            "run",
            "--with",
            "xlsxwriter",
            "--with",
            "beautifulsoup4",
            str(script),
            str(html_path),
            str(output_dir),
        ],
        check=True,
    )

    workbook = load_workbook(output_dir / "2026-06_交易明細.xlsx", read_only=False)
    assert workbook["統計"]["A1"].value == "LINE Bank 2026年06月 統計摘要"
    deposit_rows = list(workbook["台幣存款記錄"].iter_rows(values_only=True))
    assert deposit_rows[0] == ("入帳日", "交易說明", "金額", "餘額")
    assert deposit_rows[1][0].date() == date(2026, 6, 1)
    card_rows = list(workbook["刷卡記錄"].iter_rows(values_only=True))
    assert card_rows[0] == (
        "消費日",
        "入帳日",
        "商戶",
        "金額",
        "外幣消費金額",
        "換匯日",
    )
    assert card_rows[1][0].date() == date(2026, 6, 2)
    assert card_rows[1][1].date() == date(2026, 6, 3)
    assert card_rows[1][2:5] == ("商店", 50, "USD 1.50")
    assert card_rows[1][5].date() == date(2026, 6, 3)
    assert card_rows[2][0].date() == date(2026, 6, 4)
    assert card_rows[2][1:] == (None, "另一商店", 25, "-", None)
    assert workbook["台幣存款記錄"]["A2"].number_format == "yyyy-mm-dd"
    assert workbook["刷卡記錄"]["A2"].number_format == "yyyy-mm-dd"
    assert workbook["刷卡記錄"]["B2"].number_format == "yyyy-mm-dd"
    assert workbook["刷卡記錄"]["F2"].number_format == "yyyy-mm-dd"
    for sheet_name in ("統計", "台幣存款記錄", "刷卡記錄"):
        sheet = workbook[sheet_name]
        assert sheet.sheet_view.zoomScale in (None, 100)
        assert all(
            10 <= sheet.column_dimensions[get_column_letter(column)].width <= 41
            for column in range(1, sheet.max_column + 1)
        )


def test_statement_without_card_table_writes_empty_card_sheet(tmp_path):
    html_path = tmp_path / "2026-01.html"
    html_path.write_text(
        """
        <html><body>
          <div>對帳單期間: 2026/01/01~2026/01/31</div>
          <section id="tab2">
            <div class="chart-section">$100</div>
            <div class="table-scroll"><table>
              <tr><th>日期</th><th>交易說明</th><th>金額</th><th>餘額</th></tr>
              <tr><td>01/01</td><td>存入</td><td>$100</td><td>$100</td></tr>
            </table></div>
          </section>
        </body></html>
        """,
        encoding="utf-8",
    )

    script = Path(__file__).parents[1] / "scripts" / "gen_linebank_xlsx.py"
    output_dir = tmp_path / "output"
    subprocess.run(
        [
            "uv",
            "run",
            "--with",
            "xlsxwriter",
            "--with",
            "beautifulsoup4",
            str(script),
            str(html_path),
            str(output_dir),
        ],
        check=True,
    )

    workbook = load_workbook(output_dir / "2026-01_交易明細.xlsx")
    assert workbook.sheetnames == ["統計", "台幣存款記錄", "刷卡記錄"]
    card_rows = list(workbook["刷卡記錄"].iter_rows(values_only=True))
    assert card_rows[0] == (
        "消費日",
        "入帳日",
        "交易說明",
        "新臺幣金額",
        "消費國家",
        "外幣消費金額",
        "換匯日",
        "支付帳戶帳號",
    )
    assert all(not any(row) for row in card_rows[1:])


def test_statement_without_deposit_table_is_rejected(tmp_path):
    html_path = tmp_path / "2026-01.html"
    html_path.write_text(
        "<html><body><div>對帳單期間: 2026/01/01~2026/01/31</div></body></html>",
        encoding="utf-8",
    )

    script = Path(__file__).parents[1] / "scripts" / "gen_linebank_xlsx.py"
    with pytest.raises(subprocess.CalledProcessError) as exc_info:
        subprocess.run(
            [
                "uv",
                "run",
                "--with",
                "xlsxwriter",
                "--with",
                "beautifulsoup4",
                str(script),
                str(html_path),
                str(tmp_path / "output"),
            ],
            check=True,
            capture_output=True,
            text=True,
        )

    assert "缺少必要的台幣存款表格" in exc_info.value.stderr
